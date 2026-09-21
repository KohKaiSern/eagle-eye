"""Tests for the DOCX reader's output contract."""

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from backend.extraction.readers.docx_reader import read_docx


class ReadDOCXTests(TestCase):
    def test_returns_paragraphs_and_tables_in_document_order(self) -> None:
        with TemporaryDirectory() as directory:
            filepath = Path(directory) / "contract.docx"
            document = Document()
            document.add_heading("Service Agreement", level=1)
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Milestone"
            table.cell(0, 1).text = "Due date"
            table.cell(1, 0).text = "Payment"
            table.cell(1, 1).text = "15 September 2026"
            document.add_paragraph("Either party may terminate with notice.")
            document.save(filepath)

            result = read_docx(str(filepath))

        self.assertEqual(
            result,
            {
                "text": (
                    "Service Agreement\n"
                    "Milestone\tDue date\n"
                    "Payment\t15 September 2026\n"
                    "Either party may terminate with notice."
                ),
                "confidence": 1.0,
            },
        )

    def test_returns_zero_confidence_when_no_text_is_found(self) -> None:
        with TemporaryDirectory() as directory:
            filepath = Path(directory) / "blank.docx"
            Document().save(filepath)

            result = read_docx(str(filepath))

        self.assertEqual(result, {"text": "", "confidence": 0.0})
